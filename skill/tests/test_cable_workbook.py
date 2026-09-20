from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "电缆工作簿.py"
SPEC = importlib.util.spec_from_file_location("cable_workbook", SCRIPT)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def record(row_id: str, kind: str, number: int, cable_id: str = ""):
    return {
        "row_id": row_id,
        "row_kind": kind,
        "cable_type": "K",
        "number": number,
        "cable_id": cable_id,
        "start": "CAB-A" if kind == "data" else "",
        "end": "CAB-B" if kind == "data" else "",
        "spec": "" if kind == "blank" else "SPEC",
        "required_cores": None if kind == "blank" else 2,
        "selected_cores": None if kind == "blank" else 4,
        "spare_cores": None if kind == "blank" else 2,
        "core_section": "" if kind == "blank" else "",
        "length": "",
        "conduit_diameter": "",
        "conduit_length": "",
        "remark": "",
        "wire_numbers": [] if kind == "blank" else ["W1", "W2"],
        "direction_confidence": "unknown" if kind == "blank" else "high",
        "status": "confirmed",
        "issue": "",
        "source_terminal_row_ids": [] if kind == "blank" else ["t1", "t2"],
    }


def sample_model():
    return {
        "schema_version": MODULE.SCHEMA_VERSION,
        "metadata": {"knowledge_sources": ["wiki/cable.md"]},
        "records": [record("K/1", "data", 1, "K1"), record("K/2", "blank", 2), record("K/3", "data", 3, "K3")],
        "issues": [],
    }


def test_cable_workbook_roundtrip_preserves_gap_row(tmp_path: Path):
    source = sample_model()
    workbook = tmp_path / "cable.xlsx"
    MODULE.export_workbook(source, workbook)
    restored = MODULE.import_workbook(workbook)
    assert restored["metadata"] == source["metadata"]
    assert [r["cable_id"] for r in restored["records"] if r["row_kind"] == "data"] == ["K1", "K3"]
    assert restored["records"][1]["row_kind"] == "blank"
    assert restored["records"][0]["selected_cores"] == 4
    assert restored["records"][0]["wire_numbers"] == ["W1", "W2"]
    assert restored["records"][0]["source_terminal_row_ids"] == ["t1", "t2"]


def test_cable_workbook_uses_english_comma_for_joined_fields(tmp_path: Path):
    workbook = tmp_path / "separator.xlsx"
    MODULE.export_workbook(sample_model(), workbook)
    from openpyxl import load_workbook
    wb = load_workbook(workbook, data_only=True)
    assert wb["K_1"]["G2"].value == "W1,W2"
    assert wb["K_1"]["C2"].value == "t1,t2"



def test_cable_model_requires_explicit_gap_rows():
    model = sample_model()
    model["records"] = [model["records"][0], model["records"][2]]
    with pytest.raises(ValueError, match="缺少编号占位"):
        MODULE.validate_model(model)


def test_cable_model_rejects_duplicate_type_number():
    model = sample_model()
    duplicate = record("K/1-copy", "data", 1, "K1-copy")
    model["records"].append(duplicate)
    with pytest.raises(ValueError, match="数字编号必须从1开始且唯一"):
        MODULE.validate_model(model)


def test_cable_workbook_uses_block_pages_and_new_columns(tmp_path: Path):
    source = {"schema_version": MODULE.SCHEMA_VERSION, "metadata": {}, "records": [], "issues": []}
    for number in range(1, 71):
        source["records"].append(record(f"K/{number}", "data" if number in {1, 28, 69, 70} else "blank", number, f"K{number}" if number in {1, 28, 69, 70} else ""))
    workbook = tmp_path / "pages.xlsx"
    MODULE.export_workbook(source, workbook)
    from openpyxl import load_workbook
    wb = load_workbook(workbook, data_only=True)
    assert [name for name in wb.sheetnames if name not in {"待确认", "_meta"}] == ["K_1", "K_2", "K_3"]
    headers = [cell.value for cell in wb["K_1"][1]]
    assert "电缆类型" not in headers and "数字编号" not in headers
    assert headers[3:15] == ["电缆号", "起点", "终点", "线号", "需用芯数", "电缆型号", "选用芯数", "每芯截面mm²", "长度(m)", "钢管直径(mm)", "钢管长度(m)", "备注"]
    assert wb["K_1"].max_row == 32
    assert wb["K_2"].max_row == 32
    assert wb["K_3"].max_row == 32
    assert wb["K_1"].sheet_format.defaultRowHeight == 12
    assert wb["K_1"].row_dimensions[1].height == 15
    assert wb["K_1"].row_dimensions[2].height == 12
    assert wb["K_1"]["G2"].alignment.horizontal == "left"
    assert wb["K_1"]["G2"].alignment.wrap_text is not True
    assert wb["K_1"].column_dimensions["G"].width >= 32
