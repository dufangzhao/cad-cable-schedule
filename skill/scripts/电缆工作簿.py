from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCHEMA_VERSION = "cable-workbook"
SHEET_ISSUES = "待确认"
SHEET_META = "_meta"
DATA_CAPACITY = 31
TYPES = ("K", "AK", "TK")
ROW_KIND = {"data", "blank"}
STATUS = {"confirmed", "candidate", "conflict", "unpaired"}
CONFIDENCE = {"high", "medium", "low", "unknown"}
SEVERITY = {"warning", "blocking"}

# A/B/C are hidden technical columns. All visible columns mirror the standard
# cable-table block order and deliberately omit cable_type/number.
RECORD_COLUMNS = [
    ("row_id", "row_id"),
    ("row_kind", "row_kind"),
    ("source_terminal_row_ids", "source_terminal_row_ids"),
    ("cable_id", "电缆号"),
    ("start", "起点"),
    ("end", "终点"),
    ("wire_numbers", "线号"),
    ("required_cores", "需用芯数"),
    ("spec", "电缆型号"),
    ("selected_cores", "选用芯数"),
    ("core_section", "每芯截面mm²"),
    ("length", "长度(m)"),
    ("conduit_diameter", "钢管直径(mm)"),
    ("conduit_length", "钢管长度(m)"),
    ("remark", "备注"),
    ("direction_confidence", "置信度"),
    ("status", "状态"),
    ("issue", "问题与待确认"),
]
LEGACY_COLUMNS = [
    ("row_id", "row_id"), ("row_kind", "row_kind"),
    ("cable_type", "电缆类型"), ("number", "数字编号"),
    ("cable_id", "电缆号"), ("start", "起点"), ("end", "终点"),
    ("spec", "型号规格"), ("required_cores", "需用芯数"),
    ("spare_cores", "备用芯数"), ("wire_numbers", "线号"),
    ("direction_confidence", "方向置信度"), ("status", "状态"),
    ("issue", "问题与待确认"), ("source_terminal_row_ids", "source_terminal_row_ids"),
]
BLANK_EMPTY_FIELDS = {
    "cable_id", "start", "end", "spec", "required_cores", "selected_cores",
    "core_section", "length", "conduit_diameter", "conduit_length", "remark",
    "wire_numbers", "issue", "source_terminal_row_ids",
}


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def list_value(value: Any) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        values = [text(item) for item in value if text(item)]
    else:
        values = [item.strip() for item in re.split(r"[,、]", text(value)) if item.strip()]
    return list(dict.fromkeys(values))


def integer(value: Any, field: str, *, allow_blank: bool = False) -> int | None:
    if value in (None, "") and allow_blank:
        return None
    if isinstance(value, bool):
        raise ValueError(f"{field} 必须是整数")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field} 必须是整数") from exc
    if number < 0:
        raise ValueError(f"{field} 不能为负数")
    return number


def load_json(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict):
        raise ValueError("模型顶层必须是对象")
    return data


def selected_cores_for(required: Any) -> int | None:
    if required in (None, ""):
        return None
    value = integer(required, "需用芯数", allow_blank=True)
    if value is None:
        return None
    for option in (4, 7, 10):
        if option >= value:
            return option
    return None


def validate_model(model: dict[str, Any]) -> dict[str, int]:
    if model.get("schema_version") != SCHEMA_VERSION:
        raise ValueError(f"schema_version 必须是 {SCHEMA_VERSION}")
    if not isinstance(model.get("metadata", {}), dict):
        raise ValueError("metadata 必须是对象")
    records = model.get("records")
    issues = model.get("issues", [])
    if not isinstance(records, list) or not isinstance(issues, list):
        raise ValueError("records 和 issues 必须是数组")
    row_ids: set[str] = set()
    numbers: dict[str, set[int]] = {kind: set() for kind in TYPES}
    for index, record in enumerate(records):
        if not isinstance(record, dict):
            raise ValueError(f"records[{index}] 必须是对象")
        row_id = text(record.get("row_id"))
        if not row_id or row_id in row_ids:
            raise ValueError(f"records[{index}].row_id 缺失或重复")
        row_ids.add(row_id)
        cable_type = text(record.get("cable_type"))
        if cable_type not in TYPES:
            raise ValueError(f"records[{index}].cable_type 无效")
        row_kind = record.get("row_kind")
        if row_kind not in ROW_KIND:
            raise ValueError(f"records[{index}].row_kind 无效")
        number = integer(record.get("number"), f"records[{index}].number")
        if number is None or number < 1 or number in numbers[cable_type]:
            raise ValueError(f"{cable_type} 数字编号必须从1开始且唯一：{number}")
        numbers[cable_type].add(number)
        if record.get("status") not in STATUS:
            raise ValueError(f"records[{index}].status 无效")
        if record.get("direction_confidence", "unknown") not in CONFIDENCE:
            raise ValueError(f"records[{index}].direction_confidence 无效")
        for key in ("required_cores", "selected_cores", "spare_cores"):
            integer(record.get(key), f"records[{index}].{key}", allow_blank=True)
        list_value(record.get("wire_numbers"))
        list_value(record.get("source_terminal_row_ids"))
        if row_kind == "data" and not text(record.get("cable_id")):
            raise ValueError(f"records[{index}].cable_id 不能为空")
        if row_kind == "blank" and any(record.get(key) not in (None, "", []) for key in BLANK_EMPTY_FIELDS):
            raise ValueError(f"records[{index}] blank 行包含业务值")
    for cable_type, present in numbers.items():
        if present and present != set(range(1, max(present) + 1)):
            missing = sorted(set(range(1, max(present) + 1)) - present)
            raise ValueError(f"{cable_type} 缺少编号占位：{missing}")
    issue_ids: set[str] = set()
    for index, issue in enumerate(issues):
        if not isinstance(issue, dict):
            raise ValueError(f"issues[{index}] 必须是对象")
        issue_id = text(issue.get("issue_id"))
        if not issue_id or issue_id in issue_ids:
            raise ValueError(f"issues[{index}].issue_id 缺失或重复")
        issue_ids.add(issue_id)
        if issue.get("severity") not in SEVERITY:
            raise ValueError(f"issues[{index}].severity 无效")
        if text(issue.get("row_id")) and text(issue["row_id"]) not in row_ids:
            raise ValueError(f"issues[{index}].row_id 不存在")
        if not text(issue.get("message")):
            raise ValueError(f"issues[{index}].message 不能为空")
    return {"records": len(records), "issues": len(issues)}


def style_sheet(sheet: Any, columns: list[tuple[str, str]]) -> None:
    fill = PatternFill("solid", fgColor="FCE4D6")
    sheet.sheet_format.defaultRowHeight = 12
    sheet.row_dimensions[1].height = 15
    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 12
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    long_text = {"start", "end", "spec", "remark", "issue"}
    for row in sheet.iter_rows(min_row=2, max_col=len(columns)):
        for cell, (key, _) in zip(row, columns):
            cell.alignment = Alignment(
                horizontal="left" if key == "wire_numbers" else "center",
                vertical="center",
                wrap_text=(key in long_text),
                shrink_to_fit=False,
            )
    base_widths = {
        "row_id": 16, "row_kind": 10, "source_terminal_row_ids": 18,
        "cable_id": 12, "start": 18, "end": 18, "wire_numbers": 32,
        "required_cores": 12, "spec": 20, "selected_cores": 12,
        "core_section": 15, "length": 12, "conduit_diameter": 16,
        "conduit_length": 16, "remark": 20, "direction_confidence": 12,
        "status": 12, "issue": 28,
    }
    for index, (key, label) in enumerate(columns, start=1):
        values = [label]
        values.extend("" if cell.value is None else str(cell.value) for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2, values_only=False).__next__())
        max_length = max((len(value) for value in values), default=0)
        if key == "wire_numbers":
            # The line-number column is intentionally allowed to become wide enough
            # for the longest joined line-number string without wrapping.
            width = max(32, max_length * 1.15 + 2)
        else:
            width = max(base_widths.get(key, 14), min(45, max_length * 1.05 + 2))
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = "D2" if columns[:3] == [("row_id", "row_id"), ("row_kind", "row_kind"), ("source_terminal_row_ids", "source_terminal_row_ids")] else "A2"
    for column in ("A", "B", "C"):
        sheet.column_dimensions[column].hidden = True
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(columns)).coordinate}"


def blank_record(cable_type: str, number: int) -> dict[str, Any]:
    return {
        "row_id": f"{cable_type}/{number}", "row_kind": "blank", "cable_type": cable_type, "number": number,
        "cable_id": "", "start": "", "end": "", "spec": "", "required_cores": None,
        "selected_cores": None, "spare_cores": None, "core_section": "", "length": "",
        "conduit_diameter": "", "conduit_length": "", "remark": "", "wire_numbers": [],
        "direction_confidence": "unknown", "status": "confirmed", "issue": "", "source_terminal_row_ids": [],
    }


def record_value(record: dict[str, Any], key: str) -> Any:
    if key == "selected_cores" and record.get(key) in (None, ""):
        return selected_cores_for(record.get("required_cores"))
    if key == "wire_numbers" or key == "source_terminal_row_ids":
        return ",".join(list_value(record.get(key)))
    return record.get(key, "")


def export_workbook(model: dict[str, Any], output: Path) -> None:
    validate_model(model)
    if output.exists():
        raise FileExistsError(f"拒绝覆盖现有文件：{output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for cable_type in TYPES:
        records = {int(item["number"]): item for item in model["records"] if item.get("cable_type") == cable_type}
        if not records:
            continue
        page_count = math.ceil(max(records) / DATA_CAPACITY)
        for page in range(1, page_count + 1):
            sheet = workbook.create_sheet(f"{cable_type}_{page}")
            sheet.append([label for _, label in RECORD_COLUMNS])
            start = (page - 1) * DATA_CAPACITY + 1
            end = page * DATA_CAPACITY
            for number in range(start, end + 1):
                record = records.get(number, blank_record(cable_type, number))
                sheet.append([record_value(record, key) for key, _ in RECORD_COLUMNS])
            style_sheet(sheet, RECORD_COLUMNS)
    issue_sheet = workbook.create_sheet(SHEET_ISSUES)
    issue_sheet.append(["issue_id", "row_id", "字段", "级别", "问题"])
    for issue in model.get("issues", []):
        issue_sheet.append([issue.get(key, "") for key in ("issue_id", "row_id", "field", "severity", "message")])
    style_sheet(issue_sheet, [("issue_id", "issue_id"), ("row_id", "row_id"), ("field", "字段"), ("severity", "级别"), ("message", "问题")])
    meta_sheet = workbook.create_sheet(SHEET_META)
    meta_sheet.append(["schema_version", SCHEMA_VERSION])
    meta_sheet.append(["metadata_json", json.dumps(model.get("metadata", {}), ensure_ascii=False, sort_keys=True)])
    meta_sheet.append(["sheet_layout", "每个 CAD 电缆表块对应一个 sheet；每个 sheet 固定 31 行数据+1 行表头"])
    meta_sheet.sheet_state = "hidden"
    workbook.save(output)


def reject_formulas(workbook: Any) -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    raise ValueError(f"不允许公式单元格：{sheet.title}!{cell.coordinate}")


def sheet_records(sheet: Any, columns: list[tuple[str, str]]) -> list[dict[str, Any]]:
    headers = [text(cell.value) for cell in sheet[1]]
    expected = [label for _, label in columns]
    if headers[: len(expected)] != expected:
        raise ValueError(f"{sheet.title} 列名或顺序不符合契约")
    result = []
    for values in sheet.iter_rows(min_row=2, max_col=len(columns), values_only=True):
        if any(value not in (None, "") for value in values):
            result.append({key: value for (key, _), value in zip(columns, values)})
    return result


def parse_sheet_title(title: str) -> tuple[str, int]:
    match = re.fullmatch(r"(K|AK|TK)_(\d+)", title)
    if not match:
        raise ValueError(f"工作表名称必须为 K_1/AK_1/TK_1：{title}")
    return match.group(1), int(match.group(2))


def import_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    if not {SHEET_ISSUES, SHEET_META}.issubset(workbook.sheetnames):
        raise ValueError("缺少 待确认 或 _meta sheet")
    reject_formulas(workbook)
    meta = {text(row[0]): row[1] for row in workbook[SHEET_META].iter_rows(values_only=True) if row and row[0]}
    if meta.get("schema_version") != SCHEMA_VERSION:
        raise ValueError(f"_meta schema_version 必须是 {SCHEMA_VERSION}")
    metadata = json.loads(meta.get("metadata_json") or "{}")
    records: list[dict[str, Any]] = []
    data_sheets = [name for name in workbook.sheetnames if name not in {SHEET_ISSUES, SHEET_META}]
    if not data_sheets:
        raise ValueError("至少需要一个 K_n/AK_n/TK_n 工作表")
    for sheet_name in data_sheets:
        cable_type, page = parse_sheet_title(sheet_name)
        values = sheet_records(workbook[sheet_name], RECORD_COLUMNS)
        for offset, record in enumerate(values):
            row_id = text(record.get("row_id"))
            number_match = re.fullmatch(r"(?:K|AK|TK)/(\d+)", row_id)
            number = int(number_match.group(1)) if number_match else (page - 1) * DATA_CAPACITY + offset + 1
            record["cable_type"] = cable_type
            record["number"] = number
            for key in ("required_cores", "selected_cores"):
                record[key] = integer(record.get(key), key, allow_blank=True)
            record["spare_cores"] = (record["selected_cores"] - record["required_cores"] if record.get("selected_cores") is not None and record.get("required_cores") is not None else None)
            for key in ("wire_numbers", "source_terminal_row_ids"):
                record[key] = list_value(record.get(key))
            for key, _ in RECORD_COLUMNS:
                if key not in {"required_cores", "selected_cores", "wire_numbers", "source_terminal_row_ids"}:
                    record[key] = text(record.get(key))
            records.append(record)
    issues = sheet_records(workbook[SHEET_ISSUES], [("issue_id", "issue_id"), ("row_id", "row_id"), ("field", "字段"), ("severity", "级别"), ("message", "问题")])
    for issue in issues:
        for key in ("issue_id", "row_id", "field", "severity", "message"):
            issue[key] = text(issue.get(key))
    model = {"schema_version": SCHEMA_VERSION, "metadata": metadata, "records": records, "issues": issues}
    validate_model(model)
    return model


def write_json_new(data: dict[str, Any], output: Path) -> None:
    if output.exists():
        raise FileExistsError(f"拒绝覆盖现有文件：{output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="导入、导出和校验与 CAD 电缆表块一一对应的 Excel")
    subs = parser.add_subparsers(dest="command", required=True)
    exp = subs.add_parser("export"); exp.add_argument("--input-json", required=True); exp.add_argument("--output", required=True)
    imp = subs.add_parser("import"); imp.add_argument("--input", required=True); imp.add_argument("--output-json", required=True)
    val = subs.add_parser("validate"); val.add_argument("--input", required=True)
    args = parser.parse_args()
    try:
        if args.command == "export":
            model = load_json(Path(args.input_json)); export_workbook(model, Path(args.output)); result = validate_model(model) | {"output": str(Path(args.output).resolve())}
        elif args.command == "import":
            model = import_workbook(Path(args.input)); write_json_new(model, Path(args.output_json)); result = validate_model(model) | {"output": str(Path(args.output_json).resolve())}
        else:
            result = validate_model(import_workbook(Path(args.input)))
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(json.dumps({"valid": False, "error": str(exc)}, ensure_ascii=False)); return 2
    print(json.dumps({"valid": True, "schema_version": SCHEMA_VERSION, **result}, ensure_ascii=False)); return 0

if __name__ == "__main__":
    raise SystemExit(main())
