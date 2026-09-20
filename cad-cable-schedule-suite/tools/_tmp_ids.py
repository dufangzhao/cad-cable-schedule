# -*- coding: utf-8 -*-
import importlib.util, pathlib, tempfile, unicodedata
from openpyxl import load_workbook
spec = importlib.util.spec_from_file_location('core', 'deliverables/offline-plugin/core/aggregate.py')
core = importlib.util.module_from_spec(spec); spec.loader.exec_module(core)
def w(v):
    return sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in str(v or ''))
ids = [f'K{i}' for i in range(1, 301)]        # 300 个电缆号 → 拼接后约 1400 字符，远超 Excel 255 上限
model = {'schema_version': 'cable-summary', 'metadata': {},
  'issues': [{'issue_id': 'I1', 'cable_ids': ['K9'], 'field': 'length', 'severity': 'warn',
              'message': '长度未填：' + '这一行说明很长很长' * 12}],
  'records': [{'row_id': 'S1', 'spec': 'KVV', 'selected_cores': '4', 'core_section': '1.5',
               'total_length': '26', 'total_length_value': 26.0, 'quantity': 300,
               'cable_ids': ids, 'length_unknown_rows': 0}],
  'spare_records': []}
out = pathlib.Path(tempfile.mkdtemp()) / 'many.xlsx'
core.export_summary(model, out)
wb = load_workbook(out)
ws = wb['明细']
cell = ws['G2']
print('电缆号单元格字符数:', len(cell.value), '显示宽度:', w(cell.value))
print('G 列宽:', ws.column_dimensions['G'].width)
print('G 列 wrap_text:', cell.alignment.wrap_text)
ws2 = wb['待确认']
print('待确认 E2 字符数:', len(ws2['E2'].value), '列宽:', ws2.column_dimensions['E'].width, 'wrap:', ws2['E2'].alignment.wrap_text)
ws3 = wb['电缆汇总']
print('汇总表列宽:', [ws3.column_dimensions[chr(64+c)].width for c in range(1,6)])
print('汇总表表头显示宽度:', [w(ws3.cell(row=1, column=c).value) for c in range(1,6)])