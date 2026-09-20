# -*- coding: utf-8 -*-
"""端到端测试：TSV/DXF → 服务 → 下载 → 文件落盘 → LISP 响应文件。

需要服务已在 --server 指定地址运行；由 tools/run_tests.py 负责启停。
"""
from __future__ import annotations

import importlib.util
import io
import json
import subprocess
import sys
import urllib.error
import urllib.request
import uuid
from pathlib import Path

import pytest
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
SERVER = "http://127.0.0.1:8100"


def post_file(filename: str, payload: bytes, fields: dict[str, str] | None = None):
    """向服务上传一个文件，返回 (状态码, JSON)。"""
    boundary = "----t" + uuid.uuid4().hex
    buf = io.BytesIO()

    def w(x):
        buf.write(x.encode("utf-8") if isinstance(x, str) else x)

    for key, value in (fields or {}).items():
        w(f"--{boundary}\r\n")
        w(f'Content-Disposition: form-data; name="{key}"\r\n\r\n{value}\r\n')
    w(f"--{boundary}\r\n")
    w(f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n\r\n')
    w(payload)
    w(f"\r\n--{boundary}--\r\n")
    req = urllib.request.Request(f"{SERVER}/api/summarize", data=buf.getvalue(),
                                 headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    try:
        with urllib.request.urlopen(req, timeout=300) as r:
            return r.status, json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        try:
            return exc.code, json.loads(exc.read().decode("utf-8"))
        except Exception:
            return exc.code, {}


def download(ticket: str) -> bytes:
    with urllib.request.urlopen(f"{SERVER}/api/download?id={ticket}", timeout=60) as r:
        return r.read()


def _load_core():
    sys.path.insert(0, str(PROJECT))
    from core import aggregate
    return aggregate


core = _load_core()


# --------------------------------------------------------------------------- #
# 测试数据：两张电缆表 + 一张负荷表（后者必须被排除）
# --------------------------------------------------------------------------- #
HDR_X = {"cable_id": 0.0, "start": 1842.6, "end": 3708.5, "wire_numbers": 9352.3,
         "required_cores": 17051.6, "spec": 19050.0, "selected_cores": 21072.0,
         "core_section": 23065.8, "length": 25082.5, "remark": 31788.8}
HDR_TEXT = {"cable_id": "电缆号", "start": "起 点", "end": "终  点", "wire_numbers": "线          号",
            "required_cores": "需用芯数", "spec": "电缆型号", "selected_cores": "选用芯数",
            "core_section": "每芯截面", "length": "长度(m)", "remark": "备     注"}
PITCH = 700.0


def _t(handle, x, y, text, height=500.0):
    return {"handle": handle, "layer": "C-4", "x": x, "y": y, "height": height, "text": text}


def make_tsv(tmp_path: Path) -> Path:
    texts = []
    n = [0]

    def nh():
        n[0] += 1
        return f"T{n[0]:05X}"

    for ox, rows in [
        (0.0, [
            {"cable_id": "K1", "spec": "KVV", "selected_cores": "4", "core_section": "1.5", "length": "13"},
            {"cable_id": "K2", "spec": "KVV", "selected_cores": "4", "core_section": "1.5", "length": "13"},
            {"cable_id": "K3", "spec": "KVV", "selected_cores": "4", "core_section": "-", "length": "-"},
            {"cable_id": "K4"},                       # 型号/芯数/截面三项全空
        ]),
        (320000.0, [
            {"cable_id": "UK1", "spec": "YJV-0.6/1kV", "selected_cores": "2", "core_section": "2.5", "length": "25"},
        ]),
    ]:
        for key, label in HDR_TEXT.items():
            texts.append(_t(nh(), ox + HDR_X[key], 0.0, label))
        for i, row in enumerate(rows):
            y = -1275.0 - i * PITCH
            for key, value in row.items():
                off = -150.0 if key in ("cable_id", "wire_numbers") else -125.0
                texts.append(_t(nh(), ox + HDR_X[key], y + off, value))
    # 负荷表：有电缆号列但没有型号/芯数/截面/长度列 → 必须排除
    for label, x in [("电缆号", 0.0), ("起 点", 1842.6), ("终  点", 3708.5), ("最大需要容量", 5907.8)]:
        texts.append(_t(nh(), 640000.0 + x, 0.0, label))
    texts.append(_t(nh(), 640000.0, -1275.0, "X1"))

    path = tmp_path / "selection.tsv"
    core.write_tsv(texts, path)
    return path


# --------------------------------------------------------------------------- #
def test_service_reachable():
    with urllib.request.urlopen(f"{SERVER}/api/health", timeout=20) as r:
        data = json.loads(r.read().decode())
    assert data["ok"] is True
    assert data["schema_version"] == "cable-summary"


def test_tsv_pipeline_returns_correct_aggregation(tmp_path):
    """上传 TSV → 下载 Excel 的完整链路。网页服务与 CAD 插件已解耦，这里直接打 HTTP。"""
    tsv = make_tsv(tmp_path)
    status, data = post_file("sel.tsv", tsv.read_bytes())
    assert status == 200, data
    meta = data["metadata"]
    # 新口径：型号+芯数+截面 三项一致即合并。K1/K2 都是 KVV/4/1.5 → 合并；
    # K3 截面为 "-" 单独一类；UK1 型号不同另起一类 → 合计 4 根、3 类。
    assert meta["total_quantity"] == 4
    assert meta["summary_row_count"] == 3
    # 总长度：KVV/4/1.5 = 13+13 = 26；YJV/2/2.5 = 25；KVV/4/- 无法相加（不计）→ 51
    assert meta["total_length"] == 51
    assert meta["recognised_tables"] == 2, "负荷表不得计入识别结果"
    assert meta["skipped_incomplete_rows"] == 1

    saved = tmp_path / "下载.xlsx"
    saved.write_bytes(download(data["ticket"]))
    wb = load_workbook(saved)
    ws = wb["电缆汇总"]
    assert [c.value for c in ws[1]] == ["电缆型号", "选用芯数", "每芯截面(mm²)", "总长度(m)", "数量"]
    # 新表结构：不再有「合计」行，表内只有表头 + 各类别行（合计数值仍在 metadata 里）
    all_rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert not any(r and r[0] == "合计" for r in all_rows), all_rows
    assert len(all_rows) == 1 + meta["summary_row_count"]
    rows = {(ws.cell(row=r, column=1).value, ws.cell(row=r, column=5).value)
            for r in range(2, ws.max_row + 1)}
    # K1/K2 合并为 2；K3 占位符一类；UK1 型号不同另起一类（并排表的行不得被邻表抢走）
    assert ("KVV", 2) in rows and ("KVV", 1) in rows and ("YJV-0.6/1kV", 1) in rows, rows
    assert sum(q for _, q in rows) == meta["total_quantity"]


def test_download_ticket_is_single_use(tmp_path):
    status, data = post_file("sel.tsv", make_tsv(tmp_path).read_bytes())
    assert status == 200
    download(data["ticket"])
    with pytest.raises(urllib.error.HTTPError) as exc:
        download(data["ticket"])
    assert exc.value.code == 404, "取件票据必须一次性，避免临时文件长期堆积"


def test_service_rejects_wrong_content(tmp_path):
    def post(filename: str, payload: bytes):
        boundary = "----t" + uuid.uuid4().hex
        buf = io.BytesIO()
        buf.write(f"--{boundary}\r\n".encode())
        buf.write(f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n\r\n'.encode())
        buf.write(payload)
        buf.write(f"\r\n--{boundary}--\r\n".encode())
        req = urllib.request.Request(f"{SERVER}/api/summarize", data=buf.getvalue(),
                                     headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                return r.status
        except urllib.error.HTTPError as exc:
            return exc.code

    assert post("drawing.dwg", b"AC1032" + b"\x00" * 64) == 415      # 扩展名不允许
    assert post("fake.dxf", b"not a dxf") == 415                       # 内容嗅探拦截
    assert post("fake.tsv", b"a\tb\n1\t2\n") == 415                    # TSV 表头不对


def test_service_fails_closed_on_empty_selection_artifact():
    """artifact 路径：空选择必须失败关闭，不得退化为整图读取。"""
    artifact = {"success": True, "read_scope": "selected",
                "selection": {"count": 0, "handles": [], "types": [], "layers": []},
                "scan": {"total": 0, "returned": 0, "truncated": False, "items": []}}
    status, data = post_file("artifact.json", json.dumps(artifact, ensure_ascii=False).encode("utf-8"))
    assert status == 422
    assert "选择集为空" in json.dumps(data, ensure_ascii=False)
